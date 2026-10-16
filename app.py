import hmac
import os
import re
import secrets
from datetime import datetime, timedelta, timezone
from decimal import Decimal, InvalidOperation
from functools import wraps
from io import BytesIO
from urllib.parse import quote_plus
from xml.sax.saxutils import escape

from dotenv import load_dotenv
from flask import Flask, abort, flash, redirect, render_template, request, send_file, session, url_for
from flask_sqlalchemy import SQLAlchemy
from sqlalchemy import func, text
from werkzeug.security import check_password_hash, generate_password_hash

load_dotenv()  # lee el archivo .env (en AWS guarda aquí las claves, nunca en el código)

app = Flask(__name__)


# ---------------------------------------------------------------------------
# CONEXIÓN A LA BASE DE DATOS  (ítem 1 de la lista de chequeo)
# Orden de prioridad:
#   1. DATABASE_URL completa (ej: mysql+pymysql://usuario:clave@host/bd)
#   2. Variables DB_HOST, DB_USER, DB_PASSWORD, DB_NAME, DB_PORT  -> MySQL en AWS RDS
#   3. SQLite local (solo para desarrollo en el computador)
# ---------------------------------------------------------------------------
def construir_uri_bd():
    url = os.getenv("DATABASE_URL")
    if url:
        return url

    host = os.getenv("DB_HOST")
    if host:
        usuario = quote_plus(os.getenv("DB_USER", "admin"))
        clave = quote_plus(os.getenv("DB_PASSWORD", ""))
        nombre = os.getenv("DB_NAME", "elixir_urbano")
        puerto = os.getenv("DB_PORT", "3306")
        return f"mysql+pymysql://{usuario}:{clave}@{host}:{puerto}/{nombre}?charset=utf8mb4"

    ruta = os.path.join(os.path.dirname(os.path.abspath(__file__)), "elixir_urbano.db")
    return f"sqlite:///{ruta}"


app.config["SQLALCHEMY_DATABASE_URI"] = construir_uri_bd()
app.config["SQLALCHEMY_TRACK_MODIFICATIONS"] = False
# RDS cierra las conexiones inactivas: pre_ping y recycle evitan el error "MySQL server has gone away"
app.config["SQLALCHEMY_ENGINE_OPTIONS"] = {"pool_pre_ping": True, "pool_recycle": 280}
app.config["SECRET_KEY"] = os.getenv("SECRET_KEY", "cambia-esta-clave-secreta")
app.config["SESSION_COOKIE_HTTPONLY"] = True
app.config["SESSION_COOKIE_SAMESITE"] = "Lax"
app.config["SESSION_COOKIE_SECURE"] = os.getenv("COOKIE_SECURE", "0") == "1"  # pon COOKIE_SECURE=1 si usas HTTPS
app.config["PERMANENT_SESSION_LIFETIME"] = timedelta(hours=8)

db = SQLAlchemy(app)

EMAIL_REGEX = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
DOCUMENTO_REGEX = re.compile(r"^[0-9A-Za-z-]{5,20}$")
TELEFONO_REGEX = re.compile(r"^[0-9+\s-]{7,20}$")


# ---------------------------------------------------------------------------
# UTILIDADES DE FORMATO
# ---------------------------------------------------------------------------
def utc_now():
    return datetime.now(timezone.utc).replace(tzinfo=None)


def hora_local(valor):
    """Colombia es UTC-5 todo el año (no usa horario de verano)."""
    return valor - timedelta(hours=5) if valor else None


def cop(valor):
    """12500 -> $12.500"""
    return "$" + "{:,.0f}".format(valor or 0).replace(",", ".")


def fecha(valor):
    return hora_local(valor).strftime("%d/%m/%Y %H:%M") if valor else ""


app.add_template_filter(cop, "cop")
app.add_template_filter(fecha, "fecha")


# ---------------------------------------------------------------------------
# MODELOS (tablas)
# ---------------------------------------------------------------------------
class Usuario(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    nombre = db.Column(db.String(120), nullable=False)
    email = db.Column(db.String(120), unique=True, nullable=False)
    password_hash = db.Column(db.String(255), nullable=False)
    creado_en = db.Column(db.DateTime, default=utc_now)

    def set_password(self, password):
        self.password_hash = generate_password_hash(password)

    def check_password(self, password):
        return check_password_hash(self.password_hash, password)


class Producto(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    nombre = db.Column(db.String(120), nullable=False)
    precio = db.Column(db.Numeric(10, 2), nullable=False)
    stock = db.Column(db.Integer, nullable=False, default=0)
    stock_minimo = db.Column(db.Integer, nullable=False, default=5)
    imagen = db.Column(db.String(200), nullable=True)
    descripcion = db.Column(db.String(255), nullable=True)
    creado_en = db.Column(db.DateTime, default=utc_now)

    @property
    def bajo_stock(self):
        return self.stock <= self.stock_minimo


class Cliente(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    nombre = db.Column(db.String(120), nullable=False)
    documento = db.Column(db.String(20), unique=True, nullable=False)
    telefono = db.Column(db.String(20), nullable=True)
    email = db.Column(db.String(120), nullable=True)
    creado_en = db.Column(db.DateTime, default=utc_now)


class Venta(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    fecha = db.Column(db.DateTime, default=utc_now, nullable=False)
    cliente_id = db.Column(db.Integer, db.ForeignKey("cliente.id", ondelete="SET NULL"), nullable=True)
    usuario_id = db.Column(db.Integer, db.ForeignKey("usuario.id", ondelete="SET NULL"), nullable=True)
    total = db.Column(db.Numeric(12, 2), nullable=False, default=0)
    estado = db.Column(db.String(10), nullable=False, default="activa")  # activa | anulada

    cliente = db.relationship("Cliente")
    usuario = db.relationship("Usuario")
    detalles = db.relationship("DetalleVenta", backref="venta", cascade="all, delete-orphan")


class DetalleVenta(db.Model):
    __tablename__ = "detalle_venta"
    id = db.Column(db.Integer, primary_key=True)
    venta_id = db.Column(db.Integer, db.ForeignKey("venta.id", ondelete="CASCADE"), nullable=False)
    producto_id = db.Column(db.Integer, db.ForeignKey("producto.id", ondelete="SET NULL"), nullable=True)
    # Se guarda el nombre y el precio del momento: el historial no cambia si luego editan o borran el producto
    nombre_producto = db.Column(db.String(120), nullable=False)
    cantidad = db.Column(db.Integer, nullable=False)
    precio_unitario = db.Column(db.Numeric(10, 2), nullable=False)
    subtotal = db.Column(db.Numeric(12, 2), nullable=False)

    producto = db.relationship("Producto")


# ---------------------------------------------------------------------------
# SESIÓN, LOGIN OBLIGATORIO Y PROTECCIÓN CSRF
# ---------------------------------------------------------------------------
def usuario_actual():
    user_id = session.get("user_id")
    if not user_id:
        return None
    return db.session.get(Usuario, user_id)


def login_required(vista):
    @wraps(vista)
    def envoltura(*args, **kwargs):
        if not usuario_actual():
            flash("Debes iniciar sesión para acceder a esta sección.", "error")
            return redirect(url_for("login"))
        return vista(*args, **kwargs)
    return envoltura


def generar_csrf():
    if "_csrf" not in session:
        session["_csrf"] = secrets.token_hex(16)
    return session["_csrf"]


app.jinja_env.globals["csrf_token"] = generar_csrf


@app.before_request
def verificar_csrf():
    """Todo formulario POST debe traer el token oculto que se generó para esta sesión."""
    if request.method == "POST":
        esperado = session.get("_csrf", "")
        recibido = request.form.get("_csrf", "")
        if not esperado or not hmac.compare_digest(esperado, recibido):
            abort(400)


@app.context_processor
def injectar_usuario():
    try:
        return {"usuario_actual": usuario_actual()}
    except Exception:
        db.session.rollback()
        return {"usuario_actual": None}


MENSAJES_ERROR = {
    400: "La solicitud no es válida o el formulario expiró. Vuelve atrás, recarga la página e intenta de nuevo.",
    404: "La página que buscas no existe.",
    405: "Esa acción no está permitida desde esta dirección.",
    500: "Ocurrió un error interno. Intenta de nuevo en unos minutos.",
}


def manejar_error(error):
    codigo = getattr(error, "code", 500)
    if codigo == 500:
        db.session.rollback()
    return render_template("error.html", codigo=codigo, mensaje=MENSAJES_ERROR.get(codigo, "Ocurrió un error.")), codigo


for _codigo in MENSAJES_ERROR:
    app.register_error_handler(_codigo, manejar_error)


# ---------------------------------------------------------------------------
# VALIDACIONES DE BACKEND (ítem 3)
# ---------------------------------------------------------------------------
def validar_producto(form):
    """Valida los datos del formulario de producto. Devuelve (datos, errores)."""
    errores = []
    nombre = form.get("nombre", "").strip()
    precio_txt = form.get("precio", "").strip()
    stock_txt = form.get("stock", "0").strip() or "0"
    minimo_txt = form.get("stock_minimo", "5").strip() or "5"
    imagen = form.get("imagen", "").strip()
    descripcion = form.get("descripcion", "").strip()

    if len(nombre) < 2 or len(nombre) > 120:
        errores.append("El nombre debe tener entre 2 y 120 caracteres.")

    precio = None
    try:
        precio = Decimal(precio_txt)
        if not precio.is_finite() or precio <= 0 or precio > Decimal("99999999.99"):
            errores.append("El precio debe ser mayor a 0.")
            precio = None
    except InvalidOperation:
        errores.append("El precio debe ser numérico.")

    stock = minimo = None
    try:
        stock = int(stock_txt)
        minimo = int(minimo_txt)
        if stock < 0 or minimo < 0:
            errores.append("El stock y el stock mínimo no pueden ser negativos.")
    except ValueError:
        errores.append("El stock y el stock mínimo deben ser números enteros.")

    if imagen and not imagen.lower().startswith(("http://", "https://")):
        errores.append("La URL de la imagen debe empezar por http:// o https://")
    if len(imagen) > 200:
        errores.append("La URL de la imagen es demasiado larga (máximo 200 caracteres).")
    if len(descripcion) > 255:
        errores.append("La descripción no puede superar 255 caracteres.")

    datos = {
        "nombre": nombre,
        "precio": precio,
        "stock": stock,
        "stock_minimo": minimo,
        "imagen": imagen or None,
        "descripcion": descripcion or None,
    }
    return datos, errores


def validar_cliente(form, cliente_id=None):
    errores = []
    nombre = form.get("nombre", "").strip()
    documento = form.get("documento", "").strip()
    telefono = form.get("telefono", "").strip()
    email = form.get("email", "").strip().lower()

    if len(nombre) < 2 or len(nombre) > 120:
        errores.append("El nombre debe tener entre 2 y 120 caracteres.")
    if not DOCUMENTO_REGEX.match(documento):
        errores.append("El documento debe tener entre 5 y 20 caracteres (letras, números o guion).")
    else:
        existente = Cliente.query.filter_by(documento=documento).first()
        if existente and existente.id != cliente_id:
            errores.append("Ya existe un cliente con ese documento.")
    if telefono and not TELEFONO_REGEX.match(telefono):
        errores.append("El teléfono solo puede tener números, espacios, + o guion (7 a 20 caracteres).")
    if email and (not EMAIL_REGEX.match(email) or len(email) > 120):
        errores.append("El correo del cliente no es válido.")

    datos = {"nombre": nombre, "documento": documento, "telefono": telefono or None, "email": email or None}
    return datos, errores


def validar_registro(nombre, email, password, confirmar):
    errores = []
    if len(nombre) < 2:
        errores.append("El nombre debe tener al menos 2 caracteres.")
    if not EMAIL_REGEX.match(email):
        errores.append("El correo electrónico no es válido.")
    if len(password) < 8:
        errores.append("La contraseña debe tener al menos 8 caracteres.")
    if password != confirmar:
        errores.append("Las contraseñas no coinciden.")
    return errores


def leer_lineas_venta(form):
    """Convierte las filas del formulario de venta en {producto_id: cantidad} y valida el formato."""
    lineas, errores = {}, []
    for pid, cant in zip(form.getlist("producto_id"), form.getlist("cantidad")):
        if not pid and not cant.strip():
            continue  # fila vacía
        try:
            pid_i, cant_i = int(pid), int(cant)
        except ValueError:
            msg = "Cada línea necesita un producto y una cantidad entera."
            if msg not in errores:
                errores.append(msg)
            continue
        if cant_i < 1:
            msg = "La cantidad de cada producto debe ser al menos 1."
            if msg not in errores:
                errores.append(msg)
            continue
        lineas[pid_i] = lineas.get(pid_i, 0) + cant_i
    if not lineas and not errores:
        errores.append("Agrega al menos un producto a la venta.")
    return lineas, errores


# ---------------------------------------------------------------------------
# RUTAS PÚBLICAS Y AUTENTICACIÓN (ítem 4)
# ---------------------------------------------------------------------------
@app.route("/")
@app.route("/index.html")
def index():
    mas_vendidos = (
        db.session.query(DetalleVenta.nombre_producto, func.sum(DetalleVenta.cantidad).label("unidades"))
        .join(Venta, Venta.id == DetalleVenta.venta_id)
        .filter(Venta.estado == "activa")
        .group_by(DetalleVenta.nombre_producto)
        .order_by(db.desc("unidades"))
        .limit(5)
        .all()
    )
    ingresos = db.session.query(func.coalesce(func.sum(Venta.total), 0)).filter(Venta.estado == "activa").scalar()
    return render_template(
        "index.html",
        total_productos=Producto.query.count(),
        total_usuarios=Usuario.query.count(),
        total_clientes=Cliente.query.count(),
        total_ventas=Venta.query.filter_by(estado="activa").count(),
        ingresos=ingresos,
        bajo_stock=Producto.query.filter(Producto.stock <= Producto.stock_minimo).count(),
        mas_vendidos=mas_vendidos,
    )


@app.route("/health")
def health():
    """Chequeo de salud: confirma que la app y la conexión a la base de datos responden (útil en AWS)."""
    try:
        db.session.execute(text("SELECT 1"))
        return {"estado": "ok", "base_de_datos": "ok"}
    except Exception:
        db.session.rollback()
        return {"estado": "error", "base_de_datos": "sin conexión"}, 503


@app.route("/registro.html", methods=["GET", "POST"])
def registro():
    if request.method == "POST":
        nombre = request.form.get("nombre", "").strip()
        email = request.form.get("email", "").strip().lower()
        password = request.form.get("password", "")
        confirmar = request.form.get("confirmar", "")

        if not nombre or not email or not password or not confirmar:
            flash("Todos los campos son obligatorios.", "error")
            return redirect(url_for("registro"))

        errores = validar_registro(nombre, email, password, confirmar)
        if errores:
            for e in errores:
                flash(e, "error")
            return redirect(url_for("registro"))

        if Usuario.query.filter_by(email=email).first():
            flash("Ese correo ya está registrado.", "error")
            return redirect(url_for("registro"))

        usuario = Usuario(nombre=nombre, email=email)
        usuario.set_password(password)
        db.session.add(usuario)
        db.session.commit()
        flash("Cuenta creada correctamente. Ahora puedes iniciar sesión.", "success")
        return redirect(url_for("login"))

    return render_template("registro.html")


@app.route("/login.html", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        usuario_ingresado = request.form.get("usuario", "").strip().lower()
        clave = request.form.get("clave", "")

        user = Usuario.query.filter_by(email=usuario_ingresado).first()
        if user and user.check_password(clave):
            session.clear()
            session.permanent = True
            session["user_id"] = user.id
            session["user_name"] = user.nombre
            flash("Bienvenido al sistema.", "success")
            return redirect(url_for("productos"))

        flash("Usuario o contraseña incorrectos.", "error")
        return redirect(url_for("login"))

    if usuario_actual():
        return redirect(url_for("productos"))
    return render_template("login.html")


@app.route("/logout", methods=["POST"])
def logout():
    session.clear()
    flash("Sesión cerrada correctamente.", "info")
    return redirect(url_for("login"))


# ---------------------------------------------------------------------------
# CRUD DE PRODUCTOS (ítem 2)
#   Create -> crear_producto     Read   -> productos / proyectofinal
#   Update -> editar_producto    Delete -> eliminar_producto
# ---------------------------------------------------------------------------
@app.route("/productos.html")
def productos():
    lista_productos = Producto.query.order_by(Producto.id.desc()).all()
    bajo_stock = [p for p in lista_productos if p.bajo_stock]
    return render_template("productos.html", productos=lista_productos, bajo_stock=bajo_stock)


@app.route("/proyectofinal.html")
def proyectofinal():
    lista_productos = Producto.query.order_by(Producto.nombre.asc()).all()
    return render_template("proyectofinal.html", productos=lista_productos)


@app.route("/proyectoelixirurbano.html")
def proyectoelixirurbano():
    lista_productos = Producto.query.order_by(Producto.nombre.asc()).all()
    return render_template("proyectoelixirurbano.html", productos=lista_productos)


@app.route("/productos/crear", methods=["POST"])
@login_required
def crear_producto():
    datos, errores = validar_producto(request.form)
    if errores:
        for e in errores:
            flash(e, "error")
        return redirect(url_for("productos"))

    db.session.add(Producto(**datos))
    db.session.commit()
    flash("Producto creado correctamente.", "success")
    return redirect(url_for("productos"))


@app.route("/productos/editar/<int:producto_id>", methods=["GET", "POST"])
@login_required
def editar_producto(producto_id):
    producto = db.get_or_404(Producto, producto_id)

    if request.method == "POST":
        datos, errores = validar_producto(request.form)
        if errores:
            for e in errores:
                flash(e, "error")
            return redirect(url_for("editar_producto", producto_id=producto_id))

        for campo, valor in datos.items():
            setattr(producto, campo, valor)
        db.session.commit()
        flash("Producto actualizado correctamente.", "success")
        return redirect(url_for("productos"))

    return render_template("editar_producto.html", producto=producto)


@app.route("/productos/eliminar/<int:producto_id>", methods=["POST"])
@login_required
def eliminar_producto(producto_id):
    producto = db.get_or_404(Producto, producto_id)
    # El historial de ventas se conserva: las líneas ya guardan el nombre y el precio del momento
    DetalleVenta.query.filter_by(producto_id=producto.id).update({"producto_id": None})
    db.session.delete(producto)
    db.session.commit()
    flash("Producto eliminado correctamente.", "info")
    return redirect(url_for("productos"))


# ---------------------------------------------------------------------------
# CRUD DE CLIENTES
# ---------------------------------------------------------------------------
@app.route("/clientes.html")
@login_required
def clientes():
    return render_template("clientes.html", clientes=Cliente.query.order_by(Cliente.nombre.asc()).all())


@app.route("/clientes/crear", methods=["POST"])
@login_required
def crear_cliente():
    datos, errores = validar_cliente(request.form)
    if errores:
        for e in errores:
            flash(e, "error")
        return redirect(url_for("clientes"))

    db.session.add(Cliente(**datos))
    db.session.commit()
    flash("Cliente creado correctamente.", "success")
    return redirect(url_for("clientes"))


@app.route("/clientes/editar/<int:cliente_id>", methods=["GET", "POST"])
@login_required
def editar_cliente(cliente_id):
    cliente = db.get_or_404(Cliente, cliente_id)

    if request.method == "POST":
        datos, errores = validar_cliente(request.form, cliente_id=cliente_id)
        if errores:
            for e in errores:
                flash(e, "error")
            return redirect(url_for("editar_cliente", cliente_id=cliente_id))

        for campo, valor in datos.items():
            setattr(cliente, campo, valor)
        db.session.commit()
        flash("Cliente actualizado correctamente.", "success")
        return redirect(url_for("clientes"))

    return render_template("editar_cliente.html", cliente=cliente)


@app.route("/clientes/eliminar/<int:cliente_id>", methods=["POST"])
@login_required
def eliminar_cliente(cliente_id):
    cliente = db.get_or_404(Cliente, cliente_id)
    Venta.query.filter_by(cliente_id=cliente.id).update({"cliente_id": None})
    db.session.delete(cliente)
    db.session.commit()
    flash("Cliente eliminado correctamente.", "info")
    return redirect(url_for("clientes"))


# ---------------------------------------------------------------------------
# MÓDULO DE VENTAS: cada venta descuenta stock; anularla lo devuelve
# ---------------------------------------------------------------------------
@app.route("/ventas.html")
@login_required
def ventas():
    lista = Venta.query.order_by(Venta.fecha.desc()).limit(300).all()
    total_activas = sum((v.total for v in lista if v.estado == "activa"), Decimal("0"))
    return render_template("ventas.html", ventas=lista, total_activas=total_activas)


@app.route("/ventas/nueva", methods=["GET", "POST"])
@login_required
def nueva_venta():
    if request.method == "POST":
        lineas, errores = leer_lineas_venta(request.form)

        cliente = None
        cid = request.form.get("cliente_id", "").strip()
        if cid:
            cliente = db.session.get(Cliente, int(cid)) if cid.isdigit() else None
            if not cliente:
                errores.append("El cliente seleccionado no existe.")

        if not errores:
            # with_for_update bloquea las filas en MySQL para que dos ventas simultáneas no vendan el mismo stock
            productos_db = {
                p.id: p for p in Producto.query.filter(Producto.id.in_(list(lineas))).with_for_update().all()
            }
            venta = Venta(cliente_id=cliente.id if cliente else None, usuario_id=session["user_id"])
            total = Decimal("0")
            for pid, cantidad in lineas.items():
                p = productos_db.get(pid)
                if not p:
                    errores.append("Uno de los productos de la venta ya no existe.")
                elif cantidad > p.stock:
                    errores.append(f"Stock insuficiente de {p.nombre}: hay {p.stock} y pediste {cantidad}.")
                else:
                    subtotal = p.precio * cantidad
                    venta.detalles.append(DetalleVenta(
                        producto_id=p.id, nombre_producto=p.nombre, cantidad=cantidad,
                        precio_unitario=p.precio, subtotal=subtotal))
                    p.stock -= cantidad
                    total += subtotal

            if not errores:
                venta.total = total
                db.session.add(venta)
                db.session.commit()
                flash(f"Venta #{venta.id} registrada. Total: {cop(total)}.", "success")
                return redirect(url_for("detalle_venta", venta_id=venta.id))
            db.session.rollback()

        for e in errores:
            flash(e, "error")
        return redirect(url_for("nueva_venta"))

    return render_template(
        "nueva_venta.html",
        productos=Producto.query.filter(Producto.stock > 0).order_by(Producto.nombre.asc()).all(),
        clientes=Cliente.query.order_by(Cliente.nombre.asc()).all(),
    )


@app.route("/ventas/<int:venta_id>")
@login_required
def detalle_venta(venta_id):
    return render_template("detalle_venta.html", venta=db.get_or_404(Venta, venta_id))


@app.route("/ventas/anular/<int:venta_id>", methods=["POST"])
@login_required
def anular_venta(venta_id):
    venta = db.get_or_404(Venta, venta_id)
    if venta.estado == "anulada":
        flash("Esa venta ya estaba anulada.", "info")
        return redirect(url_for("detalle_venta", venta_id=venta_id))

    for d in venta.detalles:
        if d.producto:  # si el producto fue eliminado no hay stock que devolver
            d.producto.stock += d.cantidad
    venta.estado = "anulada"
    db.session.commit()
    flash(f"Venta #{venta.id} anulada. El stock fue devuelto al inventario.", "info")
    return redirect(url_for("detalle_venta", venta_id=venta_id))


# ---------------------------------------------------------------------------
# REPORTES EN EXCEL Y PDF (ítem 5)
# Cada reporte se describe como un "bloque" de datos y los mismos datos sirven para ambos formatos.
# ---------------------------------------------------------------------------
def _bloque_inventario():
    filas, resaltar = [], set()
    for i, p in enumerate(Producto.query.order_by(Producto.nombre.asc()).all()):
        precio = float(p.precio)
        filas.append([p.id, p.nombre, precio, p.stock, p.stock_minimo, precio * p.stock,
                      "BAJO STOCK" if p.bajo_stock else "OK"])
        if p.bajo_stock:
            resaltar.add(i)
    return {
        "hoja": "Inventario", "titulo": "Inventario de productos",
        "encabezados": ["ID", "Producto", "Precio", "Stock", "Stock mínimo", "Valor en inventario", "Estado"],
        "filas": filas, "resaltar": resaltar, "moneda": [2, 5],
        "totales": {3: sum(f[3] for f in filas), 5: sum(f[5] for f in filas)},
        "anchos": [6, 32, 14, 10, 14, 22, 14], "anchos_cm": [1.1, 6, 2.3, 1.6, 1.8, 2.8, 2.6], "derecha": [2, 3, 4, 5],
    }


def _bloque_ventas():
    filas, resaltar, total_activas = [], set(), 0.0
    for i, v in enumerate(Venta.query.order_by(Venta.fecha.desc()).all()):
        total = float(v.total)
        filas.append([v.id, fecha(v.fecha), v.cliente.nombre if v.cliente else "Sin cliente",
                      v.usuario.nombre if v.usuario else "-", total, v.estado.upper()])
        if v.estado == "anulada":
            resaltar.add(i)
        else:
            total_activas += total
    return {
        "hoja": "Ventas", "titulo": "Historial de ventas (el total solo suma ventas activas)",
        "encabezados": ["N°", "Fecha", "Cliente", "Registrada por", "Total", "Estado"],
        "filas": filas, "resaltar": resaltar, "moneda": [4], "totales": {4: total_activas},
        "anchos": [8, 18, 30, 26, 16, 12], "anchos_cm": [1.3, 3.2, 4.6, 4, 2.6, 2.3], "derecha": [4],
    }


def _bloque_mas_vendidos():
    consulta = (
        db.session.query(
            DetalleVenta.nombre_producto,
            func.sum(DetalleVenta.cantidad).label("unidades"),
            func.sum(DetalleVenta.subtotal).label("ingresos"),
        )
        .join(Venta, Venta.id == DetalleVenta.venta_id)
        .filter(Venta.estado == "activa")
        .group_by(DetalleVenta.nombre_producto)
        .order_by(db.desc("unidades"))
        .limit(10)
        .all()
    )
    filas = [[i, n, int(u), float(ing)] for i, (n, u, ing) in enumerate(consulta, start=1)]
    return {
        "hoja": "Más vendidos", "titulo": "Medicamentos de mayor rotación (Top 10)",
        "encabezados": ["#", "Producto", "Unidades vendidas", "Ingresos"],
        "filas": filas, "resaltar": set(), "moneda": [3], "totales": {},
        "anchos": [6, 34, 20, 18], "anchos_cm": [1.2, 8, 4, 4], "derecha": [2, 3],
    }


def _responder_excel(bloques, nombre_archivo):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    rojo = PatternFill("solid", fgColor="FED7D7")
    for b in bloques:
        ws = wb.create_sheet(b["hoja"])
        ws.append(["Droguería Elixir Urbano - " + b["titulo"]])
        ws["A1"].font = Font(bold=True, size=14)
        ws.append(["Generado: " + fecha(utc_now())])
        ws.append([])
        ws.append(b["encabezados"])
        for celda in ws[4]:
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = PatternFill("solid", fgColor="2C7A7B")
            celda.alignment = Alignment(horizontal="center")
        for i, fila in enumerate(b["filas"]):
            ws.append(fila)
            if i in b["resaltar"]:
                for celda in ws[ws.max_row]:
                    celda.fill = rojo
        if b["totales"]:
            fila_total = [""] * len(b["encabezados"])
            fila_total[1] = "TOTAL"
            for col, valor in b["totales"].items():
                fila_total[col] = valor
            ws.append(fila_total)
            for celda in ws[ws.max_row]:
                celda.font = Font(bold=True)
        for fila in ws.iter_rows(min_row=5, max_row=ws.max_row):
            for col in b["moneda"]:
                fila[col].number_format = '"$"#,##0'
        for i, ancho in enumerate(b["anchos"], start=1):
            ws.column_dimensions[get_column_letter(i)].width = ancho

    salida = BytesIO()
    wb.save(salida)
    salida.seek(0)
    return send_file(salida, as_attachment=True, download_name=nombre_archivo,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


def _responder_pdf(titulo, bloques, nombre_archivo):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    estilos = getSampleStyleSheet()
    salida = BytesIO()
    doc = SimpleDocTemplate(salida, pagesize=A4, leftMargin=1.5 * cm, rightMargin=1.5 * cm,
                            topMargin=1.5 * cm, bottomMargin=1.5 * cm, title=titulo + " - Elixir Urbano")
    elementos = [
        Paragraph("Droguería Elixir Urbano", estilos["Title"]),
        Paragraph(escape(titulo), estilos["Heading2"]),
        Paragraph(f"Generado: {fecha(utc_now())} &nbsp;|&nbsp; Usuario: {escape(usuario_actual().nombre)}", estilos["Normal"]),
        Spacer(1, 0.4 * cm),
    ]

    for b in bloques:
        if len(bloques) > 1:
            elementos += [Paragraph(escape(b["titulo"]), estilos["Heading3"])]
        datos = [b["encabezados"]]
        for fila in b["filas"]:
            datos.append([cop(v) if i in b["moneda"] else str(v)[:34] for i, v in enumerate(fila)])
        if b["totales"]:
            fila_total = [""] * len(b["encabezados"])
            fila_total[1] = "TOTAL"
            for col, valor in b["totales"].items():
                fila_total[col] = cop(valor) if col in b["moneda"] else str(valor)
            datos.append(fila_total)

        tabla = Table(datos, repeatRows=1, colWidths=[w * cm for w in b["anchos_cm"]])
        estilo = [
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2c7a7b")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cbd5e0")),
        ]
        for col in b["derecha"]:
            estilo.append(("ALIGN", (col, 1), (col, -1), "RIGHT"))
        if b["totales"]:
            estilo.append(("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"))
        for i in b["resaltar"]:
            estilo.append(("BACKGROUND", (0, i + 1), (-1, i + 1), colors.HexColor("#fed7d7")))
        tabla.setStyle(TableStyle(estilo))
        elementos += [tabla, Spacer(1, 0.6 * cm)]

    doc.build(elementos)
    salida.seek(0)
    return send_file(salida, as_attachment=True, download_name=nombre_archivo, mimetype="application/pdf")


@app.route("/reportes/excel")
@login_required
def reporte_excel():
    return _responder_excel([_bloque_inventario()], "reporte_inventario.xlsx")


@app.route("/reportes/pdf")
@login_required
def reporte_pdf():
    return _responder_pdf("Reporte de inventario", [_bloque_inventario()], "reporte_inventario.pdf")


@app.route("/reportes/ventas/excel")
@login_required
def reporte_ventas_excel():
    return _responder_excel([_bloque_ventas(), _bloque_mas_vendidos()], "reporte_ventas.xlsx")


@app.route("/reportes/ventas/pdf")
@login_required
def reporte_ventas_pdf():
    return _responder_pdf("Reporte de ventas", [_bloque_ventas(), _bloque_mas_vendidos()], "reporte_ventas.pdf")


# ---------------------------------------------------------------------------
# DATOS INICIALES
# ---------------------------------------------------------------------------
def sembrar_datos():
    if Producto.query.count() == 0:
        productos_iniciales = [
            {"nombre": "Paracetamol 500mg", "precio": 5000, "stock": 40, "stock_minimo": 10,
             "imagen": "https://placehold.co/300x180?text=Paracetamol", "descripcion": "Analgésico y antipirético."},
            {"nombre": "Ibuprofeno 400mg", "precio": 7000, "stock": 25, "stock_minimo": 10,
             "imagen": "https://placehold.co/300x180?text=Ibuprofeno", "descripcion": "Antiinflamatorio."},
            {"nombre": "Alcohol Antiséptico 70%", "precio": 4500, "stock": 8, "stock_minimo": 10,
             "imagen": "https://placehold.co/300x180?text=Alcohol+70", "descripcion": "Uso externo."},
            {"nombre": "Aspirina 100mg", "precio": 6000, "stock": 30, "stock_minimo": 10,
             "imagen": "https://placehold.co/300x180?text=Aspirina", "descripcion": "Analgésico."},
            {"nombre": "Acetaminofén", "precio": 2000, "stock": 60, "stock_minimo": 15,
             "imagen": "https://placehold.co/300x180?text=Acetaminofen", "descripcion": "Medicamento de alta rotación."},
        ]
        for p in productos_iniciales:
            db.session.add(Producto(**p))
        db.session.commit()

    if Cliente.query.count() == 0:
        db.session.add(Cliente(nombre="María Gómez", documento="1010202030", telefono="3001234567", email="maria@example.com"))
        db.session.add(Cliente(nombre="Carlos Ruiz", documento="79123456", telefono="3109876543"))
        db.session.commit()


with app.app_context():
    db.create_all()
    sembrar_datos()

if __name__ == "__main__":
    app.run(debug=True)
